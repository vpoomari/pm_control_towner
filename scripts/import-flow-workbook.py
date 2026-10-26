#!/usr/bin/env python3
"""Load the FLOW planning workbook into PMCT's SQLite project model.

Default is a read-only dry-run.  Pass --apply to write.  The workbook is only
read; --apply snapshots the database first, then applies one SQLite transaction.

Requires Python 3 + openpyxl.  Uses direct SQLite writes intentionally so this
bulk plan load does not invoke or rerun the schedule engine.  Planning Lab's
read APIs calculate diagnostics from the stored task/dependency graph on demand.
"""

from __future__ import annotations

import argparse
import collections
import datetime as dt
import json
import os
import re
import shutil
import sqlite3
import sys
import uuid
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = Path(
    r"F:\PM Control Imported Data\202601020 - Project Flow - Planning Q4 2026 - MA.xlsx"
)
DEFAULT_DB = ROOT / "db" / "custom.db"
PROJECT_CODE = "PRJ-CZR-FLOW"
PORTFOLIO_CODE = "PTF-CZR-FLOW"
PROGRAM_CODE = "PRG-CZR-FLOW"
NOW = dt.datetime.now(dt.timezone.utc).replace(microsecond=0)
NOW_ISO = NOW.isoformat().replace("+00:00", "Z")


def new_id() -> str:
    # Prisma's cuid fields are opaque TEXT IDs; UUID strings are valid and unique.
    return uuid.uuid4().hex


def clean(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def iso(value) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.replace(tzinfo=dt.timezone.utc).isoformat().replace("+00:00", "Z")
    if isinstance(value, dt.date):
        return dt.datetime.combine(value, dt.time(), dt.timezone.utc).isoformat().replace("+00:00", "Z")
    parsed = dt.datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=dt.timezone.utc)
    return parsed.astimezone(dt.timezone.utc).isoformat().replace("+00:00", "Z")


def day(value) -> dt.date | None:
    val = iso(value)
    return dt.date.fromisoformat(val[:10]) if val else None


def pct(value) -> float:
    try:
        n = float(value or 0)
    except (TypeError, ValueError):
        return 0.0
    if 0 <= n <= 1 and n != 0:
        n *= 100
    return min(100.0, max(0.0, n))


def task_status(value, progress) -> str:
    raw = (clean(value) or "NOT_STARTED").strip().upper().replace(" ", "_")
    if raw in {"NOT_STARTED", "IN_PROGRESS", "COMPLETED", "BLOCKED", "CANCELLED"}:
        return raw
    aliases = {"NOT_STARTED": "NOT_STARTED", "IN PROGRESS": "IN_PROGRESS", "DONE": "COMPLETED"}
    if raw in aliases:
        return aliases[raw]
    return "COMPLETED" if pct(progress) >= 100 else "NOT_STARTED"


def priority(value) -> str:
    raw = (clean(value) or "MEDIUM").upper()
    return raw if raw in {"LOW", "MEDIUM", "HIGH", "CRITICAL"} else "MEDIUM"


def sheet_records(sheet):
    iterator = sheet.iter_rows(values_only=True)
    headers = next(iterator)
    out = []
    for values in iterator:
        if not values or values[0] is None:
            continue
        out.append({str(h).strip(): v for h, v in zip(headers, values) if h is not None})
    return out


def read_source(path: Path):
    wb = load_workbook(path, data_only=True, read_only=True)
    plan = sheet_records(wb["Plan"])
    raid = sheet_records(wb["RAID"])
    critical_path = sheet_records(wb["Critical path"])
    wb.close()
    return plan, raid, critical_path


def validate(plan, raid, critical_path, users):
    problems: list[str] = []
    user_by_name = {u["name"].strip().casefold(): u for u in users}
    codes = [clean(r.get("ID")) for r in plan]
    if None in codes:
        problems.append("Plan: a row has no activity ID")
    duplicates = [code for code, count in collections.Counter(codes).items() if count > 1]
    if duplicates:
        problems.append(f"Plan: duplicate IDs: {', '.join(duplicates[:10])}")
    code_set = set(codes)
    unknown_owners = sorted({clean(r.get("Owner")) for r in plan if clean(r.get("Owner")) and clean(r.get("Owner")).casefold() not in user_by_name})
    if unknown_owners:
        problems.append(f"Plan: owner names have no account match: {', '.join(unknown_owners)}")

    edges = []
    external = []
    bad_dates = []
    for row in plan:
        start = day(row.get("Start date"))
        finish = day(row.get("End date"))
        if start and finish and finish < start:
            bad_dates.append(clean(row.get("ID")))
        for ref in (clean(row.get("Dependencies")) or "").split(","):
            ref = ref.strip()
            if not ref:
                continue
            if ref in code_set:
                edges.append((ref, clean(row.get("ID"))))
            else:
                external.append((clean(row.get("ID")), ref))
    if bad_dates:
        problems.append(f"Plan: end before start for {', '.join(bad_dates[:10])}")

    # Detect internal dependency cycles before touching the DB.
    graph = {c: [] for c in code_set}
    for predecessor, successor in edges:
        graph[predecessor].append(successor)
    visiting, visited = set(), set()

    def visit(node):
        if node in visiting:
            return False
        if node in visited:
            return True
        visiting.add(node)
        for nxt in graph[node]:
            if not visit(nxt):
                return False
        visiting.remove(node)
        visited.add(node)
        return True

    for code in graph:
        if not visit(code):
            problems.append(f"Plan: dependency cycle detected at {code}")
            break

    raid_ids = [str(clean(r.get("RAID ID"))) for r in raid]
    if len(set((clean(r.get("Type")), str(clean(r.get("RAID ID")))) for r in raid)) != len(raid):
        problems.append("RAID: duplicate Type + RAID ID pair")
    expected_types = {"Risk", "Issue", "Assumption", "Dependency"}
    unexpected = sorted({clean(r.get("Type")) for r in raid} - expected_types)
    if unexpected:
        problems.append(f"RAID: unsupported record types: {', '.join(map(str, unexpected))}")
    if not critical_path:
        problems.append("Critical path sheet is empty")
    return {
        "problems": problems,
        "edges": edges,
        "external": external,
        "owners": user_by_name,
        "dateRange": (
            min((day(r.get("Start date")) for r in plan if day(r.get("Start date"))), default=None),
            max((day(r.get("End date")) for r in plan if day(r.get("End date"))), default=None),
        ),
        "blankStatusPriority": sum(not clean(r.get("Status")) or not clean(r.get("Priority")) for r in plan),
        "wbsPairs": {(clean(r.get("WS")), clean(r.get("Sub"))) for r in plan},
        "workstreams": sorted({clean(r.get("WS")) for r in plan if clean(r.get("WS"))}),
        "raidTypes": dict(collections.Counter(clean(r.get("Type")) for r in raid)),
        "externalCount": len(external),
        "edgeCount": len(edges),
        "raidIdCount": len(raid_ids),
    }


def table_count(conn, table, project_id=None):
    if project_id is None:
        return conn.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0]
    return conn.execute(f'SELECT COUNT(*) FROM "{table}" WHERE projectId = ?', (project_id,)).fetchone()[0]


def backup_database(db_path: Path):
    backup_dir = ROOT / "backups"
    backup_dir.mkdir(exist_ok=True)
    stamp = NOW.strftime("%Y%m%d-%H%M%S")
    dest = backup_dir / f"custom-before-flow-workbook-{stamp}.db"
    src = sqlite3.connect(str(db_path), timeout=30)
    target = sqlite3.connect(str(dest))
    try:
        src.backup(target)
    finally:
        target.close()
        src.close()
    return dest


def upsert_project_hierarchy(conn, plan, users):
    first_start = min(day(r.get("Start date")) for r in plan if day(r.get("Start date")))
    last_finish = max(day(r.get("End date")) for r in plan if day(r.get("End date")))
    start_iso = iso(first_start)
    finish_iso = iso(last_finish)
    portfolio = conn.execute('SELECT id FROM Portfolio WHERE code = ?', (PORTFOLIO_CODE,)).fetchone()
    portfolio_id = portfolio[0] if portfolio else new_id()
    conn.execute(
        """INSERT INTO Portfolio (id, code, name, description, status, strategicObjective,
             currency, startDate, endDate, createdAt, updatedAt)
           VALUES (?, ?, ?, ?, 'ACTIVE', ?, 'USD', ?, ?, ?, ?)
           ON CONFLICT(code) DO UPDATE SET name=excluded.name, description=excluded.description,
             strategicObjective=excluded.strategicObjective, startDate=excluded.startDate,
             endDate=excluded.endDate, updatedAt=excluded.updatedAt""",
        (portfolio_id, PORTFOLIO_CODE, "FLOW Enterprise Transformation", "Portfolio for the FLOW enterprise transformation planning workbook.", "Integrated enterprise transformation delivery", start_iso, finish_iso, NOW_ISO, NOW_ISO),
    )
    portfolio_id = conn.execute('SELECT id FROM Portfolio WHERE code = ?', (PORTFOLIO_CODE,)).fetchone()[0]

    program = conn.execute('SELECT id FROM Program WHERE code = ?', (PROGRAM_CODE,)).fetchone()
    program_id = program[0] if program else new_id()
    conn.execute(
        """INSERT INTO Program (id, code, name, description, portfolioId, status, startDate, endDate, createdAt, updatedAt)
           VALUES (?, ?, ?, ?, ?, 'ACTIVE', ?, ?, ?, ?)
           ON CONFLICT(code) DO UPDATE SET name=excluded.name, description=excluded.description,
             portfolioId=excluded.portfolioId, startDate=excluded.startDate, endDate=excluded.endDate,
             updatedAt=excluded.updatedAt""",
        (program_id, PROGRAM_CODE, "FLOW Enterprise Transformation", "Integrated programme plan sourced from the Q4 2026 FLOW planning workbook.", portfolio_id, start_iso, finish_iso, NOW_ISO, NOW_ISO),
    )
    program_id = conn.execute('SELECT id FROM Program WHERE code = ?', (PROGRAM_CODE,)).fetchone()[0]

    project = conn.execute('SELECT id FROM Project WHERE code = ?', (PROJECT_CODE,)).fetchone()
    project_id = project[0] if project else new_id()
    conn.execute(
        """INSERT INTO Project (id, code, name, description, portfolioId, programId, ownerId,
             managerId, status, priority, phase, methodology, riskLevel, startDate, endDate,
             statusDate, objectives, createdAt, updatedAt)
           VALUES (?, ?, ?, ?, ?, ?, NULL, NULL, 'ACTIVE', 'MEDIUM', 'PLANNING', 'HYBRID', 'MEDIUM', ?, ?, ?, ?, ?, ?)
           ON CONFLICT(code) DO UPDATE SET name=excluded.name, description=excluded.description,
             portfolioId=excluded.portfolioId, programId=excluded.programId,
             ownerId=NULL, managerId=NULL,
             startDate=excluded.startDate, endDate=excluded.endDate,
             statusDate=excluded.statusDate, objectives=excluded.objectives, updatedAt=excluded.updatedAt""",
        (project_id, PROJECT_CODE, "FLOW - Enterprise Transformation", "Integrated enterprise transformation plan loaded from 202601020 - Project Flow - Planning Q4 2026 - MA.xlsx.", portfolio_id, program_id, start_iso, finish_iso, NOW_ISO, "Plan activities, dependencies, workstream structure, RAID and critical path as supplied by the planning workbook.", NOW_ISO, NOW_ISO),
    )
    project_id = conn.execute('SELECT id FROM Project WHERE code = ?', (PROJECT_CODE,)).fetchone()[0]
    return portfolio_id, program_id, project_id


def upsert_wbs(conn, project_id, plan):
    by_pair = {}
    for row in plan:
        ws, sub = clean(row.get("WS")), clean(row.get("Sub"))
        if ws and sub:
            by_pair.setdefault(ws, set()).add(sub)
    root_ids = {}
    sub_ids = {}
    root_codes = {}
    for order, ws in enumerate(sorted(by_pair), 1):
        slug = re.sub(r"[^A-Z0-9]+", "-", ws.upper()).strip("-")
        code = f"FLOW-WS-{slug}"
        old = conn.execute('SELECT id FROM WBSNode WHERE projectId=? AND code=?', (project_id, code)).fetchone()
        node_id = old[0] if old else new_id()
        conn.execute(
            """INSERT INTO WBSNode (id, projectId, parentId, code, name, description, nodeType,
                 level, orderIndex, createdAt, updatedAt)
               VALUES (?, ?, NULL, ?, ?, ?, 'SUMMARY', 1, ?, ?, ?)
               ON CONFLICT(id) DO UPDATE SET name=excluded.name, description=excluded.description,
                 nodeType='SUMMARY', level=1, orderIndex=excluded.orderIndex, updatedAt=excluded.updatedAt""",
            (node_id, project_id, code, ws, f"FLOW workbook workstream: {ws}", order, NOW_ISO, NOW_ISO),
        )
        root_ids[ws] = node_id
        root_codes[ws] = code
        for sub_order, sub in enumerate(sorted(by_pair[ws]), 1):
            sub_slug = re.sub(r"[^A-Z0-9]+", "-", sub.upper()).strip("-")
            sub_code = f"{code}-{sub_slug}"
            old_sub = conn.execute('SELECT id FROM WBSNode WHERE projectId=? AND code=?', (project_id, sub_code)).fetchone()
            sub_id = old_sub[0] if old_sub else new_id()
            conn.execute(
                """INSERT INTO WBSNode (id, projectId, parentId, code, name, description, nodeType,
                     level, orderIndex, createdAt, updatedAt)
                   VALUES (?, ?, ?, ?, ?, ?, 'WORK_PACKAGE', 2, ?, ?, ?)
                   ON CONFLICT(id) DO UPDATE SET parentId=excluded.parentId, name=excluded.name,
                     description=excluded.description, nodeType='WORK_PACKAGE', level=2,
                     orderIndex=excluded.orderIndex, updatedAt=excluded.updatedAt""",
                (sub_id, project_id, node_id, sub_code, sub, f"{ws} / {sub}", sub_order, NOW_ISO, NOW_ISO),
            )
            sub_ids[(ws, sub)] = sub_id
    return sub_ids, len(root_ids) + len(sub_ids)


def upsert_tasks(conn, project_id, plan, sub_ids, owner_by_name, external_by_task):
    task_ids = {}
    for order, row in enumerate(plan, 1):
        code = clean(row.get("ID"))
        old = conn.execute('SELECT id FROM Task WHERE projectId=? AND code=?', (project_id, code)).fetchone()
        task_id = old[0] if old else new_id()
        task_ids[code] = task_id

    for order, row in enumerate(plan, 1):
        code = clean(row.get("ID"))
        task_id = task_ids[code]
        ws, sub = clean(row.get("WS")), clean(row.get("Sub"))
        start = iso(row.get("Start date"))
        finish = iso(row.get("End date"))
        start_day, finish_day = day(row.get("Start date")), day(row.get("End date"))
        duration = float((finish_day - start_day).days + 1) if start_day and finish_day else 1.0
        progress = pct(row.get("% completed"))
        status = task_status(row.get("Status"), progress)
        task_priority = priority(row.get("Priority"))
        owner_name = clean(row.get("Owner"))
        owner = owner_by_name.get(owner_name.casefold()) if owner_name else None
        comments = clean(row.get("Comments"))
        notes = []
        if clean(row.get("Phase")):
            notes.append(f"Workbook phase: {clean(row.get('Phase'))}")
        if clean(row.get("Date basis")):
            notes.append(f"Date basis: {clean(row.get('Date basis'))}")
        external_refs = external_by_task.get(code, [])
        if external_refs:
            notes.append("External dependencies (not task-linked): " + ", ".join(external_refs))
        note_text = "\n".join(notes) or None
        remaining = 0.0 if progress >= 100 else 0.0  # workbook contains no effort estimate
        conn.execute(
            """INSERT INTO Task (id, projectId, wbsId, parentId, code, name, description, status,
                 priority, criticality, startDate, endDate, durationDays, progress, plannedHours,
                 actualHours, remainingHours, plannedCost, actualCost, assigneeId, isCritical,
                 isSummary, totalFloat, freeFloat, constraintType, notes, createdAt, updatedAt,
                 backlogOrder)
               VALUES (?, ?, ?, NULL, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0, 0, ?, 0, 0, ?, ?, 0,
                 0, 0, 'ASAP', ?, ?, ?, ?)
               ON CONFLICT(id) DO UPDATE SET wbsId=excluded.wbsId, name=excluded.name,
                 description=excluded.description, status=excluded.status, priority=excluded.priority,
                 criticality=excluded.criticality, startDate=excluded.startDate, endDate=excluded.endDate,
                 durationDays=excluded.durationDays, progress=excluded.progress,
                 remainingHours=excluded.remainingHours, assigneeId=excluded.assigneeId,
                 isCritical=excluded.isCritical, notes=excluded.notes,
                 backlogOrder=excluded.backlogOrder, updatedAt=excluded.updatedAt""",
            (task_id, project_id, sub_ids.get((ws, sub)), code, clean(row.get("Activity")) or code,
             comments, status, task_priority, task_priority, start, finish, duration, progress,
             remaining, owner["id"] if owner else None, False, note_text,
             NOW_ISO, NOW_ISO, order),
        )
    return task_ids


def upsert_dependencies(conn, project_id, plan, task_ids):
    inserted = updated = 0
    for row in plan:
        successor_code = clean(row.get("ID"))
        for predecessor_code in (clean(row.get("Dependencies")) or "").split(","):
            predecessor_code = predecessor_code.strip()
            if not predecessor_code or predecessor_code not in task_ids:
                continue
            predecessor_id, successor_id = task_ids[predecessor_code], task_ids[successor_code]
            existing = conn.execute(
                "SELECT id FROM Dependency WHERE projectId=? AND predecessorId=? AND successorId=?",
                (project_id, predecessor_id, successor_id),
            ).fetchone()
            if existing:
                conn.execute("UPDATE Dependency SET depType='FS', lagDays=0, isExternal=0, externalRef=NULL WHERE id=?", (existing[0],))
                updated += 1
            else:
                conn.execute(
                    "INSERT INTO Dependency (id, projectId, predecessorId, successorId, depType, lagDays, isExternal, createdAt) VALUES (?, ?, ?, ?, 'FS', 0, 0, ?)",
                    (new_id(), project_id, predecessor_id, successor_id, NOW_ISO),
                )
                inserted += 1
    return inserted, updated


def import_raid(conn, project_id, raid):
    totals = collections.Counter()
    for row in raid:
        raw_id = str(clean(row.get("RAID ID")))
        kind = clean(row.get("Type"))
        ws, title = clean(row.get("Workstream")), clean(row.get("Title")) or f"RAID {raw_id}"
        owner, due = clean(row.get("Owner")), iso(row.get("Due date"))
        description = clean(row.get("Description"))
        mitigation = clean(row.get("Mitigation / Action"))
        escalation = clean(row.get("Escalation / Decision needed"))
        status = (clean(row.get("Status")) or "OPEN").upper().replace(" ", "_")
        severity = (clean(row.get("Severity")) or "MEDIUM").upper()
        if kind == "Risk":
            code = f"R-{int(raw_id):03d}"
            prob = {"LOW": 2, "MEDIUM": 3, "HIGH": 4, "VERY HIGH": 5}.get((clean(row.get("Likelihood")) or "MEDIUM").upper(), 3)
            impact = {"LOW": 2, "MEDIUM": 3, "HIGH": 4, "VERY HIGH": 5}.get((clean(row.get("Impact")) or "MEDIUM").upper(), 3)
            existing = conn.execute('SELECT id FROM Risk WHERE projectId=? AND code=?', (project_id, code)).fetchone()
            rid = existing[0] if existing else new_id()
            conn.execute(
                """INSERT INTO Risk (id, projectId, code, title, description, category, probability,
                     impact, score, severity, status, ownerName, responseStrategy, mitigation,
                     contingency, escalationLevel, identifiedAt, dueDate, updatedAt)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'MITIGATE', ?, NULL, ?, ?, ?, ?)
                   ON CONFLICT(id) DO UPDATE SET title=excluded.title, description=excluded.description,
                     category=excluded.category, probability=excluded.probability, impact=excluded.impact,
                     score=excluded.score, severity=excluded.severity, status=excluded.status,
                     ownerName=excluded.ownerName, mitigation=excluded.mitigation,
                     escalationLevel=excluded.escalationLevel, dueDate=excluded.dueDate,
                     updatedAt=excluded.updatedAt""",
                (rid, project_id, code, title, description, (ws or "GENERAL").upper().replace("/", "_"), prob, impact, prob * impact, severity, status, owner, mitigation, "PMO" if escalation else "NONE", iso(row.get("Last updated")) or NOW_ISO, due, NOW_ISO),
            )
            totals["risks"] += 1
        elif kind in {"Issue", "Dependency"}:
            code = f"{'I' if kind == 'Issue' else 'D'}-{int(raw_id):03d}"
            issue_description = (f"Workstream: {ws}\nRecord type: {kind}\n\n" if ws else f"Record type: {kind}\n\n") + (description or "")
            impact_text = clean(row.get("Impact"))
            priority_value = "HIGH" if severity in {"HIGH", "CRITICAL"} else "MEDIUM"
            existing = conn.execute('SELECT id FROM Issue WHERE projectId=? AND code=?', (project_id, code)).fetchone()
            issue_id = existing[0] if existing else new_id()
            conn.execute(
                """INSERT INTO Issue (id, projectId, code, title, description, category, priority,
                     severity, status, ownerName, raisedBy, impact, resolution, escalationLevel,
                     raisedAt, dueDate, resolvedAt, updatedAt)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, NULL, ?, ?, ?, ?, ?, NULL, ?)
                   ON CONFLICT(id) DO UPDATE SET title=excluded.title, description=excluded.description,
                     category=excluded.category, priority=excluded.priority, severity=excluded.severity,
                     status=excluded.status, ownerName=excluded.ownerName, impact=excluded.impact,
                     resolution=excluded.resolution, escalationLevel=excluded.escalationLevel,
                     dueDate=excluded.dueDate, updatedAt=excluded.updatedAt""",
                (issue_id, project_id, code, title, issue_description, "DEPENDENCY" if kind == "Dependency" else (ws or "GENERAL").upper().replace("/", "_"), priority_value, severity, "OPEN" if status in {"OPEN", "TO_VALIDATE"} else status, owner, impact_text, mitigation or escalation, "PMO" if escalation else "NONE", iso(row.get("Last updated")) or NOW_ISO, due, NOW_ISO),
            )
            totals["issues"] += 1
        elif kind == "Assumption":
            code = f"A-{int(raw_id):03d}"
            assumption_status = "TO_VALIDATE" if "VALIDATE" in status else ("VALID" if status in {"OPEN", "VALID"} else status)
            rationale = f"Workstream: {ws}" if ws else None
            existing = conn.execute('SELECT id FROM Assumption WHERE projectId=? AND code=?', (project_id, code)).fetchone()
            assumption_id = existing[0] if existing else new_id()
            conn.execute(
                """INSERT INTO Assumption (id, projectId, code, description, rationale, impactIfFalse,
                     status, ownerName, validationDate, dueDate, createdAt)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, NULL, ?, ?)
                   ON CONFLICT(id) DO UPDATE SET description=excluded.description, rationale=excluded.rationale,
                     impactIfFalse=excluded.impactIfFalse, status=excluded.status, ownerName=excluded.ownerName,
                     dueDate=excluded.dueDate""",
                (assumption_id, project_id, code, title, rationale, description, assumption_status, owner, due, NOW_ISO),
            )
            totals["assumptions"] += 1
    return dict(totals)


def import_milestones(conn, project_id, critical_path):
    count = 0
    for row in critical_path:
        try:
            seq = int(row.get("Seq"))
        except (TypeError, ValueError):
            continue
        code = f"FLOW-CP-{seq:02d}"
        existing = conn.execute('SELECT id FROM Milestone WHERE projectId=? AND code=?', (project_id, code)).fetchone()
        milestone_id = existing[0] if existing else new_id()
        stage, outcome, lead = clean(row.get("Stage")), clean(row.get("Outcome")), clean(row.get("Lead"))
        desc = f"Critical-path stage outcome: {outcome or 'Not supplied'}\nLead: {lead or 'Not supplied'}"
        due = iso(row.get("End"))
        conn.execute(
            """INSERT INTO Milestone (id, projectId, code, name, description, dueDate, baselineDate,
                 status, isCritical, createdAt)
               VALUES (?, ?, ?, ?, ?, ?, NULL, 'PENDING', 1, ?)
               ON CONFLICT(id) DO UPDATE SET name=excluded.name, description=excluded.description,
                 dueDate=excluded.dueDate, status='PENDING', isCritical=1""",
            (milestone_id, project_id, code, stage or code, desc, due, NOW_ISO),
        )
        count += 1
    return count


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--db", type=Path, default=DEFAULT_DB)
    parser.add_argument("--apply", action="store_true", help="Apply after full preflight; default is dry-run")
    args = parser.parse_args()
    if not args.workbook.is_file():
        raise SystemExit(f"Workbook not found: {args.workbook}")
    if not args.db.is_file():
        raise SystemExit(f"SQLite database not found: {args.db}")

    plan, raid, critical_path = read_source(args.workbook)
    conn = sqlite3.connect(str(args.db), timeout=45)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys=ON")
    users = conn.execute('SELECT id, name, email FROM User').fetchall()
    report = validate(plan, raid, critical_path, users)

    print(f"Mode: {'APPLY' if args.apply else 'DRY RUN'}")
    print(f"Workbook: {args.workbook}")
    print(f"Database: {args.db}")
    print(f"Source rows: tasks={len(plan)} WBS groups={len(report['workstreams'])} WBS pairs={len(report['wbsPairs'])} internal dependencies={report['edgeCount']} external references={report['externalCount']}")
    print(f"RAID={len(raid)} ({json.dumps(report['raidTypes'], sort_keys=True)}) critical-path stages={len(critical_path)}")
    print(f"Date range: {report['dateRange'][0]} to {report['dateRange'][1]}")
    print(f"Blank status/priority rows (will default NOT_STARTED/MEDIUM): {report['blankStatusPriority']}")
    project = conn.execute('SELECT id FROM Project WHERE code=?', (PROJECT_CODE,)).fetchone()
    if project:
        print(f"Existing target project rows: tasks={table_count(conn, 'Task', project['id'])} wbs={table_count(conn, 'WBSNode', project['id'])} dependencies={table_count(conn, 'Dependency', project['id'])} risks={table_count(conn, 'Risk', project['id'])} issues={table_count(conn, 'Issue', project['id'])} assumptions={table_count(conn, 'Assumption', project['id'])} milestones={table_count(conn, 'Milestone', project['id'])}")
    else:
        print("Existing target project: absent; hierarchy will be created and linked")
    for message in report["problems"]:
        print(f"PRE-FLIGHT ERROR: {message}")
    if report["problems"]:
        conn.close()
        raise SystemExit(2)
    print("Owner linkage: all activity owners match existing User accounts")
    if report["external"]:
        print("External dependency refs: " + "; ".join(f"{task} -> {ref}" for task, ref in report["external"]))
    if not args.apply:
        print("Dry run only: database unchanged. Use --apply to snapshot and load.")
        conn.close()
        return

    snapshot = backup_database(args.db)
    print(f"Pre-import DB snapshot: {snapshot}")
    try:
        conn.execute("BEGIN IMMEDIATE")
        # Re-fetch owners inside the write transaction for stable linkage.
        users = conn.execute('SELECT id, name, email FROM User').fetchall()
        _, _, project_id = upsert_project_hierarchy(conn, plan, users)
        sub_ids, wbs_count = upsert_wbs(conn, project_id, plan)
        owner_by_name = {u["name"].strip().casefold(): u for u in users}
        task_ids = upsert_tasks(conn, project_id, plan, sub_ids, owner_by_name, {
            task: [ref for task_code, ref in report["external"] if task_code == task]
            for task in {task for task, _ in report["external"]}
        })
        dep_inserted, dep_updated = upsert_dependencies(conn, project_id, plan, task_ids)
        raid_counts = import_raid(conn, project_id, raid)
        milestones = import_milestones(conn, project_id, critical_path)

        # Append a normal audit event without logging source content or secrets.
        actor = conn.execute('SELECT id, name FROM User WHERE email=? COLLATE NOCASE', ('pmo@pmct.io',)).fetchone()
        audit_id = new_id()
        after = {
            "workbook": args.workbook.name,
            "projectCode": PROJECT_CODE,
            "portfolioCode": PORTFOLIO_CODE,
            "programCode": PROGRAM_CODE,
            "tasks": len(plan), "wbsNodes": wbs_count,
            "dependenciesInserted": dep_inserted, "dependenciesUpdated": dep_updated,
            **raid_counts, "milestones": milestones,
            "externalRefsPreservedInTaskNotes": report["externalCount"],
        }
        conn.execute(
            """INSERT INTO AuditEvent (id, userId, userName, role, action, entityType,
                 entityId, entityName, afterJson, context, severity, createdAt)
               VALUES (?, ?, ?, 'PMO_ADMIN', 'IMPORT', 'FLOW_WORKBOOK', ?, ?, ?, ?, 'INFO', ?)""",
            (audit_id, actor["id"] if actor else None, actor["name"] if actor else "Workbook Import", project_id,
             args.workbook.name, json.dumps(after, ensure_ascii=False), "Direct SQLite import; no schedule recalculation invoked.", NOW_ISO),
        )
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()

    # Reopen read-only and verify the persisted graph/links after commit.
    verify = sqlite3.connect(f"file:{args.db.as_posix()}?mode=ro", uri=True)
    verify.row_factory = sqlite3.Row
    project_id = verify.execute('SELECT id FROM Project WHERE code=?', (PROJECT_CODE,)).fetchone()[0]
    checks = {
        "portfolioProgramProjectLinked": verify.execute(
            "SELECT COUNT(*) FROM Project p JOIN Program g ON p.programId=g.id JOIN Portfolio f ON p.portfolioId=f.id WHERE p.id=? AND g.portfolioId=f.id",
            (project_id,),
        ).fetchone()[0] == 1,
        "tasks": table_count(verify, "Task", project_id),
        "tasksWithValidWbs": verify.execute('SELECT COUNT(*) FROM Task t JOIN WBSNode w ON t.wbsId=w.id WHERE t.projectId=?', (project_id,)).fetchone()[0],
        "tasksWithOwner": verify.execute('SELECT COUNT(*) FROM Task WHERE projectId=? AND assigneeId IS NOT NULL', (project_id,)).fetchone()[0],
        "allTaskDependencyLinksValid": verify.execute('SELECT COUNT(*) FROM Dependency d JOIN Task p ON d.predecessorId=p.id JOIN Task s ON d.successorId=s.id WHERE d.projectId=? AND p.projectId=? AND s.projectId=?', (project_id, project_id, project_id)).fetchone()[0],
        "dependencies": table_count(verify, "Dependency", project_id),
        "risks": table_count(verify, "Risk", project_id),
        "issues": table_count(verify, "Issue", project_id),
        "assumptions": table_count(verify, "Assumption", project_id),
        "milestones": table_count(verify, "Milestone", project_id),
        "auditRows": verify.execute("SELECT COUNT(*) FROM AuditEvent WHERE entityType='FLOW_WORKBOOK' AND entityId=?", (project_id,)).fetchone()[0],
        "dateRange": [verify.execute('SELECT MIN(startDate), MAX(endDate) FROM Task WHERE projectId=?', (project_id,)).fetchone()[i] for i in range(2)],
    }
    verify.close()
    print("APPLIED AND VERIFIED:")
    print(json.dumps(checks, indent=2))
    if checks["tasks"] != len(plan) or checks["tasksWithValidWbs"] != len(plan) or checks["dependencies"] != report["edgeCount"]:
        raise SystemExit("Post-import verification failed: task/WBS/dependency counts differ from source")
    if not checks["portfolioProgramProjectLinked"] or checks["allTaskDependencyLinksValid"] != checks["dependencies"]:
        raise SystemExit("Post-import verification failed: hierarchy or dependency linkage is broken")
    if checks["auditRows"] < 1:
        raise SystemExit("Post-import verification failed: missing audit row")


if __name__ == "__main__":
    main()
