#!/usr/bin/env python3
"""
PM CONTROL TOWER — training document exporter
=============================================
Converts docs/training/*.md into:
  * docs/training/pdf/  — print-styled PDFs  (Markdown -> HTML -> headless Chrome)
  * docs/training/docx/ — Word documents     (python-docx)

Usage:
  python scripts/export-training-docs.py

Requirements: python-docx, markdown, and Google Chrome or Microsoft Edge.
The script auto-detects the browser. Re-run after editing any training .md.
"""

from __future__ import annotations

import html
import re
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "training"
PDF_DIR = SRC / "pdf"
DOCX_DIR = SRC / "docx"

BRAND_NAVY = "#0b1f3a"
BRAND_CYAN = "#0e7490"
BRAND_LIGHT = "#f0f5fa"

BROWSERS = [
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
]


def find_browser() -> str:
    for b in BROWSERS:
        if Path(b).exists():
            return b
    sys.exit("No Chrome/Edge found — cannot render PDFs.")


# ---------------------------------------------------------------- PDF path --

def md_to_html(text: str, title: str) -> str:
    import markdown

    body = markdown.markdown(
        text,
        extensions=["tables", "fenced_code", "sane_lists", "nl2br"],
        output_format="html5",
    )
    css = f"""
    @page {{ size: A4; margin: 18mm 15mm 18mm 15mm; }}
    * {{ box-sizing: border-box; }}
    body {{
      font-family: 'Segoe UI', 'Helvetica Neue', Arial, sans-serif;
      color: #1e293b; font-size: 10.5pt; line-height: 1.55; margin: 0;
    }}
    h1 {{
      color: {BRAND_NAVY}; font-size: 20pt; border-bottom: 3px solid {BRAND_CYAN};
      padding-bottom: 8px; margin: 0 0 14px 0;
    }}
    h2 {{
      color: {BRAND_NAVY}; font-size: 14pt; margin: 22px 0 8px 0;
      border-left: 4px solid {BRAND_CYAN}; padding-left: 8px;
      page-break-after: avoid;
    }}
    h3 {{ color: #334155; font-size: 11.5pt; margin: 16px 0 6px 0; page-break-after: avoid; }}
    p {{ margin: 6px 0; }}
    table {{
      border-collapse: collapse; width: 100%; margin: 10px 0; font-size: 9pt;
      page-break-inside: avoid;
    }}
    th {{
      background: {BRAND_NAVY}; color: #fff; text-align: left;
      padding: 6px 8px; border: 1px solid #cbd5e1;
    }}
    td {{ padding: 5px 8px; border: 1px solid #dbe3ec; vertical-align: top; }}
    tr:nth-child(even) td {{ background: {BRAND_LIGHT}; }}
    code {{
      font-family: Consolas, 'Courier New', monospace; font-size: 9pt;
      background: #eef2f7; border: 1px solid #dbe3ec; border-radius: 3px;
      padding: 1px 4px; color: #0e7490;
    }}
    pre {{
      background: #0d1b2f; color: #d7e3f4; padding: 10px 12px; border-radius: 6px;
      font-family: Consolas, 'Courier New', monospace; font-size: 8.5pt;
      overflow-x: hidden; white-space: pre-wrap; page-break-inside: avoid;
    }}
    pre code {{ background: none; border: none; color: inherit; padding: 0; }}
    blockquote {{
      border-left: 4px solid {BRAND_CYAN}; background: {BRAND_LIGHT};
      margin: 10px 0; padding: 8px 12px; color: #334155; font-style: italic;
      page-break-inside: avoid;
    }}
    ul, ol {{ margin: 6px 0 6px 20px; padding-left: 12px; }}
    li {{ margin: 3px 0; }}
    hr {{ border: none; border-top: 1px solid #cbd5e1; margin: 16px 0; }}
    a {{ color: {BRAND_CYAN}; text-decoration: none; }}
    strong {{ color: {BRAND_NAVY}; }}
    """
    return f"""<!DOCTYPE html><html><head><meta charset="utf-8">
    <title>{html.escape(title)}</title><style>{css}</style></head>
    <body>{body}</body></html>"""


def export_pdfs(files: list[Path]) -> None:
    browser = find_browser()
    PDF_DIR.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        for md in files:
            title = md.stem.split("-", 1)[-1].replace("-", " ").title() if md.stem[0].isdigit() else md.stem
            page = tmp / f"{md.stem}.html"
            page.write_text(md_to_html(md.read_text(encoding="utf-8"), title), encoding="utf-8")
            out = PDF_DIR / f"{md.stem}.pdf"
            cmd = [
                browser, "--headless=new", "--disable-gpu", "--no-sandbox",
                "--no-pdf-header-footer", f"--print-to-pdf={out}",
                page.resolve().as_uri(),
            ]
            r = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
            if not out.exists() or out.stat().st_size < 1024:
                sys.exit(f"PDF failed for {md.name}: {r.stderr[-400:]}")
            print(f"  pdf  {out.name}  {out.stat().st_size // 1024} KB")


# --------------------------------------------------------------- DOCX path --

INLINE = re.compile(
    r"(?P<code>`[^`]+`)"
    r"|(?P<bold>\*\*[^*]+\*\*)"
    r"|(?P<italic>\*[^*\n]+\*)"
    r"|(?P<link>\[[^\]]+\]\([^)]+\))"
)

def add_inline(paragraph, text: str, base_code: bool = False) -> None:
    """Add runs to a paragraph, honouring inline markdown."""
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos:m.start()])
            run.font.name = "Segoe UI"; run.font.size = None if not base_code else None
        tok = m.group()
        if m.group("code"):
            run = paragraph.add_run(tok[1:-1])
            run.font.name = "Consolas"; run.font.size = None
            run._element.rPr.rFonts.set(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "Consolas")
            from docx.shared import Pt
            run.font.size = Pt(9)
            run.font.color.rgb = __import__("docx").shared.RGBColor(0x0E, 0x74, 0x90)
        elif m.group("bold"):
            run = paragraph.add_run(tok[2:-2]); run.bold = True
        elif m.group("italic"):
            run = paragraph.add_run(tok[1:-1]); run.italic = True
        elif m.group("link"):
            lm = re.match(r"\[([^\]]+)\]\(([^)]+)\)", tok)
            run = paragraph.add_run(f"{lm.group(1)} ({lm.group(2)})")
            run.font.color.rgb = __import__("docx").shared.RGBColor(0x0E, 0x74, 0x90)
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def export_docx(files: list[Path]) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Cm

    DOCX_DIR.mkdir(parents=True, exist_ok=True)
    navy = RGBColor(0x0B, 0x1F, 0x3A)

    for md in files:
        doc = Document()
        # page margins
        for s in doc.sections:
            s.top_margin = s.bottom_margin = Cm(2)
            s.left_margin = s.right_margin = Cm(2.2)
        # base style
        st = doc.styles["Normal"]
        st.font.name = "Segoe UI"; st.font.size = Pt(10.5)
        for lvl, size in ((1, 20), (2, 14), (3, 11.5)):
            hs = doc.styles[f"Heading {lvl}"]
            hs.font.name = "Segoe UI"; hs.font.size = Pt(size); hs.font.color.rgb = navy

        lines = md.read_text(encoding="utf-8").splitlines()
        i, n = 0, len(lines)
        while i < n:
            line = lines[i]
            stripped = line.strip()

            if stripped.startswith("```"):                       # code block
                i += 1
                buf = []
                while i < n and not lines[i].strip().startswith("```"):
                    buf.append(lines[i]); i += 1
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.4)
                run = p.add_run("\n".join(buf))
                run.font.name = "Consolas"; run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(0xD7, 0xE3, 0xF4)
                shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "0D1B2F")
                p._p.get_or_add_pPr().append(shd)
                i += 1; continue

            if stripped.startswith("|"):                          # table
                block = []
                while i < n and lines[i].strip().startswith("|"):
                    block.append(lines[i].strip()); i += 1
                rows = [[c.strip() for c in r.strip("|").split("|")] for r in block
                        if not re.match(r"^\|[\s:|-]+\|$", r)]
                if rows:
                    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
                    t.style = "Table Grid"
                    for ri, row in enumerate(rows):
                        for ci, cell in enumerate(row[: len(rows[0])]):
                            cp = t.rows[ri].cells[ci].paragraphs[0]
                            add_inline(cp, cell)
                            if ri == 0:
                                for r in cp.runs: r.bold = True
                    doc.add_paragraph()
                continue

            m_h = re.match(r"^(#{1,4})\s+(.*)$", stripped)
            if m_h:                                               # headings
                lvl = min(len(m_h.group(1)), 3)
                h = doc.add_heading(level=lvl)
                add_inline(h, m_h.group(2))
                i += 1; continue

            if stripped.startswith(("---", "===")) and set(stripped) <= set("-="):
                p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(4)
                pBdr = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
                bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "6")
                bot.set(qn("w:color"), "CBD5E1"); pBdr.append(bot)
                p._p.get_or_add_pPr().append(pBdr)
                i += 1; continue

            if stripped.startswith(">"):                           # blockquote
                buf = []
                while i < n and lines[i].strip().startswith(">"):
                    buf.append(lines[i].strip().lstrip(">").strip()); i += 1
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                run = p.add_run(" ".join(x for x in buf if x))
                run.italic = True; run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
                pPr = p._p.get_or_add_pPr()
                pBdr = OxmlElement("w:pBdr"); left = OxmlElement("w:left")
                left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "18")
                left.set(qn("w:color"), "0E7490"); left.set(qn("w:space"), "6")
                pBdr.append(left); pPr.append(pBdr)
                continue

            m_li = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
            if m_li:                                              # lists
                depth = len(m_li.group(1)) // 2
                numbered = m_li.group(2)[0].isdigit()
                style = "List Number" if numbered else "List Bullet"
                if depth > 0: style += " 2" if depth == 1 else " 3"
                try:
                    p = doc.add_paragraph(style=style)
                except KeyError:
                    p = doc.add_paragraph(style="List Bullet")
                add_inline(p, m_li.group(3))
                i += 1; continue

            if not stripped:                                      # blank
                i += 1; continue

            p = doc.add_paragraph()                               # plain para
            add_inline(p, stripped)
            i += 1

        first = lines[0].lstrip("# ").strip() if lines and lines[0].startswith("# ") else md.stem
        doc.core_properties.title = first
        doc.core_properties.author = "PM Control Tower — Training"
        out = DOCX_DIR / f"{md.stem}.docx"
        doc.save(out)
        print(f"  docx {out.name}  {out.stat().st_size // 1024} KB")


def main() -> None:
    files = sorted(SRC.glob("*.md"))
    if not files:
        sys.exit(f"No markdown files in {SRC}")
    print(f"Exporting {len(files)} training documents…")
    export_pdfs(files)
    export_docx(files)
    print(f"Done.\n  PDF : {PDF_DIR}\n  DOCX: {DOCX_DIR}")


if __name__ == "__main__":
    main()
